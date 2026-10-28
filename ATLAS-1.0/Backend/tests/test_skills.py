"""Skill-Sandbox: neue Fähigkeiten entwickeln, testen, freigeben, nutzen, zurücksetzen.

Echt: generierter Python-Code wird geprüft und in einem eigenen, begrenzten Prozess getestet;
PowerPoint-Dateien werden wirklich erzeugt und mit python-pptx nachgelesen; Freigaben und Versionen
in SQLite. Simuliert: das Sprachmodell, das den Code schreibt (FakeOllama).
"""
import contextlib
import json
import sqlite3
import tempfile
import unittest
from pathlib import Path

from atlas import llm
from atlas.approvals import Approvals
from atlas.skills import SkillLab, check_code
from atlas.tools import Environment
from fake_ollama import FakeOllama, answer

try:
    import pptx  # noqa: F401
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False

PPTX_CODE = '''
from pptx import Presentation
from pathlib import Path

def run(args, output_dir):
    title = args.get("title", "Präsentation")
    slides = args.get("slides") or [{"title": title, "text": ""}]
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[0])
    first.shapes.title.text = title
    for item in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = item.get("title", "")
        slide.placeholders[1].text = item.get("text", "")
    name = args.get("filename", "praesentation.pptx")
    deck.save(str(Path(output_dir) / name))
    return {"ok": True, "message": f"Präsentation mit {len(slides) + 1} Folien erstellt.", "files": [name]}
'''
PPTX_TESTS = '''
import os, unittest
from pathlib import Path
from pptx import Presentation
import skill

class T(unittest.TestCase):
    def test_creates_real_deck(self):
        out = os.environ["ATLAS_SKILL_OUTPUT"]
        result = skill.run({"title": "Quartal", "slides": [{"title": "Umsatz", "text": "42000 Euro"}]}, out)
        self.assertTrue(result["ok"])
        deck = Presentation(str(Path(out) / "praesentation.pptx"))
        self.assertEqual(len(deck.slides), 2)
        self.assertEqual(deck.slides[1].shapes.title.text, "Umsatz")
'''


def spec(code=PPTX_CODE, tests=PPTX_TESTS, name='create_presentation'):
    return answer(json.dumps({'name': name, 'description': 'Erstellt eine PowerPoint-Präsentation (.pptx).',
                              'parameters': {'type': 'object', 'properties': {
                                  'title': {'type': 'string'}, 'slides': {'type': 'array', 'items': {'type': 'object'}},
                                  'filename': {'type': 'string'}}, 'required': ['title']},
                              'requirements': ['pptx'], 'code': code, 'tests': tests}))


class SkillTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name).resolve()
        self.ollama = FakeOllama()
        self.addCleanup(self.ollama.close)
        self.approvals = Approvals(self.connect)
        self.installed = []
        self.lab = SkillLab(self.connect, self.root / 'skills', self.approvals, llm.Ollama(base=self.ollama.url),
                            installer=lambda package: self.installed.append(package) or f'{package} installiert.')
        self.home = self.root / 'home'
        (self.home / 'Documents').mkdir(parents=True)
        self.env = Environment(home=self.home, roots=[self.home], is_mac=False)

    @contextlib.contextmanager
    def connect(self):
        db = sqlite3.connect(self.root / 'atlas.db')
        try:
            with db:
                yield db
        finally:
            db.close()

    def develop(self):
        return self.lab.develop('Erstelle PowerPoint-Präsentationen', 'qwen3:4b')

    # -- Sicherheitsprüfung ---------------------------------------------------------------------
    def test_static_check_blocks_dangerous_code(self):
        bad = {
            'import subprocess\ndef run(a, o): subprocess.run(["rm", "-rf", "/"])': 'Import nicht erlaubt',
            'import socket': 'Import nicht erlaubt',
            'def run(a, o): eval("1")': 'Verboten: eval',
            'def run(a, o): ().__class__.__bases__': 'Verbotener Zugriff',
            'import os\ndef run(a, o): os.system("ls")': 'Verbotene Operation: .system',
            'from pathlib import Path\ndef run(a, o): Path("/x").unlink()': 'Verbotene Operation: .unlink',
            'import shutil': 'Import nicht erlaubt',
            'def run(a, o): __import__("os")': 'Verboten: __import__',
        }
        for code, expected in bad.items():
            with self.subTest(code=code[:40]):
                self.assertTrue(any(expected in p for p in check_code(code)), check_code(code))
        self.assertEqual(check_code(PPTX_CODE), [])

    def test_dangerous_skill_is_never_executed(self):
        marker = self.root / 'gehackt.txt'
        evil = f'import subprocess\ndef run(args, output_dir):\n    subprocess.run(["touch", "{marker}"])\n    return {{"ok": True}}\n'
        self.ollama.script = [spec(code=evil)] * 3
        result = self.develop()
        self.assertEqual(result['status'], 'fehlgeschlagen')
        self.assertIn('Import nicht erlaubt', result['error'])
        self.assertFalse(marker.exists())
        self.assertEqual(self.approvals.pending(), [])

    def test_network_is_blocked_inside_skills(self):
        code = ('import json\nfrom urllib import request\ndef run(args, output_dir):\n'
                '    request.urlopen("https://example.org", timeout=3)\n    return {"ok": True}\n')
        tests = 'import os, unittest, skill\nclass T(unittest.TestCase):\n    def test(self):\n        skill.run({}, os.environ["ATLAS_SKILL_OUTPUT"])\n'
        self.assertTrue(check_code(code))                    # urllib ist schon statisch verboten …
        from atlas.skills import run_harness
        folder = self.root / 'netz'
        folder.mkdir()
        (folder / 'skill.py').write_text(code.replace('from urllib import request', 'import socket as request\nrequest.urlopen = lambda *a, **k: request.create_connection(("example.org", 80))'))
        (folder / 'test_skill.py').write_text(tests)
        result = run_harness('test', folder, self.root / 'out')     # … und zur Laufzeit zusätzlich gesperrt
        self.assertFalse(result['ok'])
        self.assertIn('Netzwerk ist für Fähigkeiten gesperrt', result['output'])

    def test_writing_outside_sandbox_fails_the_tests(self):
        target = self.root / 'ausserhalb.txt'
        code = f'def run(args, output_dir):\n    open("{target}", "w").write("x")\n    return {{"ok": True, "files": []}}\n'
        tests = 'import os, unittest, skill\nclass T(unittest.TestCase):\n    def test(self):\n        skill.run({}, os.environ["ATLAS_SKILL_OUTPUT"])\n'
        self.ollama.script = [spec(code=code, tests=tests, name='schreib_test')]
        result = self.lab.develop('x', 'qwen3:4b', max_attempts=1)
        self.assertEqual(result['status'], 'fehlgeschlagen')
        self.assertIn('außerhalb', result['error'])
        self.assertFalse(target.exists())

    # -- Entwickeln, Freigeben, Nutzen, Zurücksetzen ---------------------------------------------
    @unittest.skipUnless(HAS_PPTX, 'python-pptx fehlt')
    def test_new_skill_is_developed_tested_approved_and_used(self):
        broken = PPTX_CODE.replace('deck.slide_layouts[1]', 'deck.slide_layouts[99]')
        self.ollama.script = [spec(code=broken), spec()]          # 1. Versuch scheitert, 2. besteht
        result = self.develop()
        self.assertEqual(result['status'], 'wartet_auf_freigabe', result)
        self.assertEqual(result['tests'], 1)
        retry = self.ollama.chat_requests()[1]['messages'][-1]['content']
        self.assertIn('Tests sind fehlgeschlagen', retry)              # Fehler wurde zurückgemeldet
        self.assertEqual(self.lab.tools(self.env), {})                 # ohne Freigabe nicht aktiv
        decided = self.approvals.decide(result['proposal'], True)
        self.assertEqual(decided['status'], 'aktiv', decided['result'])
        tools = self.lab.tools(self.env)
        self.assertIn('skill_create_presentation', tools)
        tool = tools['skill_create_presentation']
        self.assertEqual(tool.risk, 'confirm')
        outcome = tool.func(title='Strategie 2027', slides=[{'title': 'Ziele', 'text': 'Wachstum'}],
                            filename='strategie.pptx', output_folder='~/Documents')
        self.assertTrue(outcome.ok and outcome.verified, outcome.output)
        from pptx import Presentation
        deck = Presentation(str(self.home / 'Documents' / 'strategie.pptx'))
        self.assertEqual(deck.slides[0].shapes.title.text, 'Strategie 2027')

    @unittest.skipUnless(HAS_PPTX, 'python-pptx fehlt')
    def test_versions_and_rollback(self):
        self.ollama.script = [spec()]
        first = self.develop()
        self.approvals.decide(first['proposal'], True)
        self.ollama.script = [spec(code=PPTX_CODE.replace('Folien erstellt', 'Folien erstellt (v2)'))]
        second = self.develop()
        self.assertEqual(second['version'], 2)
        self.approvals.decide(second['proposal'], True)
        statuses = {(s['version'], s['status']) for s in self.lab.list() if s['name'] == 'create_presentation'}
        self.assertEqual(statuses, {(1, 'archiviert'), (2, 'aktiv')})
        self.assertIn('v1 wieder aktiv', self.lab.rollback('create_presentation'))
        self.assertIn('[Selbst entwickelt, v1]', self.lab.tools(self.env)['skill_create_presentation'].description)

    def test_rejected_skill_stays_inactive(self):
        if not HAS_PPTX:
            self.skipTest('python-pptx fehlt')
        self.ollama.script = [spec()]
        result = self.develop()
        self.assertEqual(self.approvals.decide(result['proposal'], False)['status'], 'abgelehnt')
        self.assertEqual(self.lab.tools(self.env), {})
        with self.assertRaises(ValueError):
            self.approvals.decide(result['proposal'], True)        # Entscheidung ist endgültig

    def test_missing_library_needs_approval_before_install(self):
        code = 'import fpdf\ndef run(args, output_dir):\n    return {"ok": True}\n'
        tests = 'import unittest, skill\nclass T(unittest.TestCase):\n    def test(self):\n        self.assertTrue(skill.run({}, ".")["ok"])\n'
        self.ollama.script = [spec(code=code, tests=tests, name='pdf_bericht')]
        try:
            import fpdf  # noqa: F401
            self.skipTest('fpdf ist hier installiert')
        except ImportError:
            pass
        result = self.lab.develop('PDF-Berichte', 'qwen3:4b')
        self.assertEqual(result['status'], 'braucht_paket')
        self.assertEqual(self.installed, [])                              # nichts ohne Freigabe installiert
        proposal = self.approvals.get(result['proposals'][0])
        self.assertEqual(proposal['detail']['package'], 'fpdf2')
        self.approvals.decide(proposal['id'], True)
        self.assertEqual(self.installed, ['fpdf2'])


if __name__ == '__main__':
    unittest.main()
