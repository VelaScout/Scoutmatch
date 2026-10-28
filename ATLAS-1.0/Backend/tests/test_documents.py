"""Dokumente: echte Dateien erstellen, lesen, ändern, umwandeln – geprüft über Format, Inhalt und Wiederöffnen.

Alles echt (python-docx, openpyxl, python-pptx, reportlab, pypdf). Kein Sprachmodell beteiligt.
"""
import tempfile
import unittest
import zipfile
from pathlib import Path

from atlas.documents import (convert_document, create_document, create_presentation, edit_document, edit_spreadsheet)
from atlas.tools import Environment, ToolError, create_spreadsheet, read_file, write_file


def pdf_text(path):
    from pypdf import PdfReader
    return ''.join(page.extract_text() for page in PdfReader(str(path)).pages)


def office_ok(path):
    """Echtes Office-Format: ZIP-Container mit [Content_Types].xml."""
    with zipfile.ZipFile(path) as archive:
        return '[Content_Types].xml' in archive.namelist() and archive.testzip() is None


class DocumentTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.home = Path(self.tmp.name).resolve()
        (self.home / 'Documents').mkdir()
        self.env = Environment(home=self.home, roots=[self.home], is_mac=False)
        self.docs = self.home / 'Documents'

    def offer(self):
        return create_document(self.env, '~/Documents/Angebot.docx', 'Angebot Webseite',
                               '# Leistungen\n- Gestaltung\n- Umsetzung\n## Preis\nGesamt: 4.800 € netto, gültig bis 30.11.2026.')

    def test_word_document_is_real_and_structured(self):
        import docx
        result = self.offer()
        path = self.docs / 'Angebot.docx'
        self.assertTrue(result.verified and office_ok(path))
        document = docx.Document(str(path))
        styles = {p.text: p.style.name for p in document.paragraphs}
        self.assertEqual(styles['Leistungen'], 'Heading 1')
        self.assertEqual(styles['Gestaltung'], 'List Bullet')
        self.assertIn('Gesamt: 4.800 € netto, gültig bis 30.11.2026.', styles)
        text = read_file(self.env, '~/Documents/Angebot.docx').output
        self.assertIn('# Leistungen', text)
        self.assertIn('- Umsetzung', text)
        with self.assertRaises(ToolError):
            self.offer()                                   # nie still überschreiben

    def test_write_file_docx_is_a_real_word_file(self):
        write_file(self.env, '~/Documents/Notiz.docx', 'Einkauf\n- Milch\n- Brot')
        self.assertTrue(office_ok(self.docs / 'Notiz.docx'))
        self.assertIn('Milch', read_file(self.env, '~/Documents/Notiz.docx').output)

    def test_read_powerpoint_and_csv(self):
        create_presentation(self.env, '~/Documents/Plan.pptx', 'Plan', [
            {'title': 'Ziele', 'bullets': ['Umsatz steigern']}, {'title': 'Schritte', 'bullets': ['Werbung']}])
        text = read_file(self.env, '~/Documents/Plan.pptx').output
        self.assertIn('[Folie 2]', text)
        self.assertIn('Werbung', text)
        (self.docs / 'daten.csv').write_text('Name;Betrag\nMiete;1200\n', encoding='utf-8')
        self.assertIn('Miete\t1200', read_file(self.env, '~/Documents/daten.csv').output)

    def test_edit_word_keeps_backup_and_refuses_missing_text(self):
        import docx
        self.offer()
        path = self.docs / 'Angebot.docx'
        before = path.read_bytes()
        with self.assertRaises(ToolError):
            edit_document(self.env, '~/Documents/Angebot.docx', [{'old': 'gibt es nicht', 'new': 'x'}])
        self.assertEqual(path.read_bytes(), before)                     # unverändert
        result = edit_document(self.env, '~/Documents/Angebot.docx', [{'old': '4.800 €', 'new': '5.200 €'}],
                               append='Zahlbar innerhalb von 14 Tagen.')
        text = '\n'.join(p.text for p in docx.Document(str(path)).paragraphs)
        self.assertIn('5.200 €', text)
        self.assertNotIn('4.800 €', text)
        self.assertIn('Zahlbar innerhalb von 14 Tagen.', text)
        saved = Path(result.data['backup'])
        self.assertTrue(saved.is_file())
        self.assertIn('4.800 €', '\n'.join(p.text for p in docx.Document(str(saved)).paragraphs))
        self.assertFalse(str(saved).startswith(str(self.docs)))         # Sicherung nicht im Nutzerordner

    def test_edit_text_file(self):
        (self.docs / 'konzept.md').write_text('# Konzept\nVersion 1\n', encoding='utf-8')
        edit_document(self.env, '~/Documents/konzept.md', [{'old': 'Version 1', 'new': 'Version 2'}], append='- Neu')
        self.assertEqual((self.docs / 'konzept.md').read_text(encoding='utf-8'), '# Konzept\nVersion 2\n- Neu\n')

    def test_edit_spreadsheet_cells_formula_and_rows(self):
        from openpyxl import load_workbook
        create_spreadsheet(self.env, '~/Documents/Budget.xlsx', columns=['Posten', 'Betrag'], rows=[['Miete', 1200], ['Strom', 90]])
        edit_spreadsheet(self.env, '~/Documents/Budget.xlsx', cells={'B3': 95}, append_rows=[['Summe', '=SUM(B2:B3)']])
        ws = load_workbook(self.docs / 'Budget.xlsx').active
        self.assertEqual(ws['B3'].value, 95)
        self.assertEqual((ws['A4'].value, ws['B4'].value), ('Summe', '=SUM(B2:B3)'))
        with self.assertRaises(ToolError):
            edit_spreadsheet(self.env, '~/Documents/Budget.xlsx', cells={'B2': 1}, sheet='Gibtsnicht')
        (self.docs / 'liste.csv').write_text('A;B\n1;2\n', encoding='utf-8')
        edit_spreadsheet(self.env, '~/Documents/liste.csv', cells={'B2': 3}, append_rows=[[4, 5]])
        self.assertEqual((self.docs / 'liste.csv').read_text(encoding='utf-8').split(), ['A;B', '1;3', '4;5'])

    def test_convert_to_pdf_from_every_supported_format(self):
        self.offer()
        create_spreadsheet(self.env, '~/Documents/Kosten.xlsx', columns=['Posten', 'Betrag'], rows=[['Miete', 1200]])
        create_presentation(self.env, '~/Documents/Vortrag.pptx', 'Vortrag', [{'title': 'Einleitung', 'bullets': ['Willkommen']}])
        (self.docs / 'notiz.md').write_text('# Notiz\nTermin am Montag', encoding='utf-8')
        for source, needle in [('Angebot.docx', '4.800'), ('Kosten.xlsx', 'Miete'),
                               ('Vortrag.pptx', 'Willkommen'), ('notiz.md', 'Montag')]:
            with self.subTest(source=source):
                result = convert_document(self.env, f'~/Documents/{source}')
                out = Path(result.data['path'])
                self.assertEqual(out.read_bytes()[:5], b'%PDF-')
                self.assertIn(needle, pdf_text(out))
        again = convert_document(self.env, '~/Documents/notiz.md')     # vorhandene PDF wird nicht überschrieben
        self.assertTrue(again.data['path'].endswith('notiz-2.pdf'))
        word = convert_document(self.env, '~/Documents/notiz.md', to='docx')
        self.assertTrue(office_ok(word.data['path']))


if __name__ == '__main__':
    unittest.main()
