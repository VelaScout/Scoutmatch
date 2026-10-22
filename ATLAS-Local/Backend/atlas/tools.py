"""Werkzeuge des Agenten mit abgestuften Berechtigungen.

Stufen:
- SAFE:    unkritisch, läuft ohne Rückfrage (lesen, neu anlegen, öffnen, suchen)
- CONFIRM: braucht eine ausdrückliche Bestätigung (überschreiben, verschieben, löschen,
           nicht-lesende Terminalbefehle)
- BLOCKED: wird nie ausgeführt (sudo, Systemänderungen, Installationen über Pipes …)

Jedes Werkzeug prüft sein Ergebnis selbst nach (``verified``), damit der Agent keine
Erfolge behauptet, die nicht tatsächlich eingetreten sind. macOS-Schutzmechanismen
(Datenschutz-Freigaben, Sandbox, Gatekeeper) werden nie umgangen.
"""
from __future__ import annotations

import csv
import html
import importlib.util
import json
import os
import platform
import re
import shlex
import shutil
import subprocess
import tempfile
import time
import urllib.parse
import urllib.request
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Callable

from .intents import calculate, normalize_expression

SAFE, CONFIRM, BLOCKED = 'safe', 'confirm', 'blocked'
MAX_READ = 60_000
MAX_OUTPUT = 12_000


class ToolError(Exception):
    """Fehler, der dem Modell als Werkzeugergebnis zurückgemeldet wird."""


@dataclass
class Result:
    ok: bool
    output: str
    verified: bool = False
    data: dict = field(default_factory=dict)

    def for_model(self) -> str:
        payload = {'ok': self.ok, 'verified': self.verified, 'output': self.output[:MAX_OUTPUT]}
        if self.data:
            payload['data'] = self.data
        return json.dumps(payload, ensure_ascii=False)


@dataclass
class Tool:
    name: str
    description: str
    parameters: dict
    func: Callable[..., Result]
    risk: str = SAFE
    category: str = 'core'
    label: str = 'Arbeitet'
    # Optional: entscheidet je Aufruf über die Stufe (z. B. Terminal: lesen vs. ändern).
    risk_for: Callable[[dict], str] | None = None
    # Kurztext für die Bestätigung.
    describe: Callable[[dict], str] | None = None

    def schema(self) -> dict:
        return {'type': 'function', 'function': {
            'name': self.name, 'description': self.description, 'parameters': self.parameters}}

    def risk_of(self, args: dict) -> str:
        return self.risk_for(args) if self.risk_for else self.risk

    def summary(self, args: dict) -> str:
        if self.describe:
            return self.describe(args)
        return f'{self.label}: ' + ', '.join(f'{k}={str(v)[:80]}' for k, v in args.items())


def obj(properties: dict, required: list[str] | None = None) -> dict:
    return {'type': 'object', 'properties': properties, 'required': required or []}


def string(description: str) -> dict:
    return {'type': 'string', 'description': description}


# ---------------------------------------------------------------------------
# Umgebung: Pfade, Prozesse, Netzwerk (in Tests austauschbar)
# ---------------------------------------------------------------------------

class Environment:
    """Alles, was die Werkzeuge vom System brauchen – in Tests ersetzbar."""

    def __init__(self, home: Path | None = None, roots: list[Path] | None = None,
                 runner: Callable[..., subprocess.CompletedProcess] | None = None,
                 fetch: Callable[[str, float], str] | None = None, is_mac: bool | None = None,
                 vision: Callable[[list[bytes], str], str | None] | None = None):
        self.home = (home or Path.home()).resolve()
        self.roots = [r.resolve() for r in (roots or [self.home])]
        self.runner = runner or subprocess.run
        self.fetch = fetch or http_get
        self.is_mac = platform.system() == 'Darwin' if is_mac is None else is_mac
        self.trash = self.home / '.Trash'
        # Lokales Bildmodell (Ollama), falls installiert: Bilder → Beschreibung/Antwort.
        self.vision = vision

    def run(self, args: list[str], timeout: float = 20, cwd: Path | None = None,
            input_text: str | None = None, extra_env: dict | None = None) -> subprocess.CompletedProcess:
        options = {'capture_output': True, 'text': True, 'timeout': timeout,
                   'cwd': str(cwd) if cwd else None, 'input': input_text}
        if extra_env:
            options['env'] = {**os.environ, **extra_env}
        return self.runner(args, **options)

    # -- Pfade ---------------------------------------------------------------
    def resolve(self, raw: str, *, must_exist: bool = False, for_write: bool = False) -> Path:
        if not isinstance(raw, str) or not raw.strip():
            raise ToolError('Pfad fehlt.')
        text = raw.strip()
        aliases = {'schreibtisch': 'Desktop', 'desktop': 'Desktop', 'dokumente': 'Documents',
                   'documents': 'Documents', 'downloads': 'Downloads', 'home': '', 'zuhause': ''}
        first, _, rest = text.replace('\\', '/').partition('/')
        if first.lower() in aliases and not text.startswith(('/', '~')):
            text = str(Path(aliases[first.lower()]) / rest) if rest else aliases[first.lower()]
        if text == '~' or text.startswith('~/'):
            path = self.home / text[2:]
        elif text.startswith('~'):
            raise ToolError('Nur der eigene Benutzerordner ist erlaubt.')
        else:
            path = Path(text)
        if not path.is_absolute():
            path = self.home / path
        path = Path(os.path.normpath(path))
        # Symlinks auflösen, damit kein Link aus dem erlaubten Bereich herausführt.
        real = path.resolve() if path.exists() else path.parent.resolve() / path.name
        if not any(real == root or root in real.parents for root in self.roots):
            raise ToolError(f'Pfad außerhalb der freigegebenen Ordner: {real}')
        relative = real.relative_to(self.home) if self.home in real.parents or real == self.home else None
        if relative is not None and relative.parts:
            top = relative.parts[0]
            if top == 'Library' or (top.startswith('.') and top not in ('.Trash',)):
                raise ToolError('Systemordner und versteckte Ordner sind gesperrt.')
            if top == '.Trash' and for_write:
                raise ToolError('Der Papierkorb kann nicht direkt beschrieben werden.')
        if must_exist and not real.exists():
            raise ToolError(f'Nicht gefunden: {real}')
        return real


def http_get(url: str, timeout: float = 15) -> str:
    request = urllib.request.Request(url, headers={
        'User-Agent': 'Mozilla/5.0 (Macintosh; Apple Silicon) Atlas/0.3', 'Accept-Language': 'de,en;q=0.8'})
    with urllib.request.urlopen(request, timeout=timeout) as response:
        kind = response.headers.get('Content-Type', '')
        if 'html' not in kind and 'text' not in kind and 'json' not in kind:
            raise ToolError(f'Kein Textinhalt ({kind or "unbekannt"}).')
        return response.read(2_000_000).decode(response.headers.get_content_charset() or 'utf-8', 'replace')


def html_to_text(raw: str) -> str:
    raw = re.sub(r'(?is)<(script|style|noscript|svg|head)[^>]*>.*?</\1>', ' ', raw)
    raw = re.sub(r'(?i)<br\s*/?>|</(p|div|li|h[1-6]|tr)>', '\n', raw)
    text = html.unescape(re.sub(r'<[^>]+>', ' ', raw))
    lines = [re.sub(r'[ \t\xa0]+', ' ', line).strip() for line in text.splitlines()]
    return '\n'.join(line for line in lines if line)


# ---------------------------------------------------------------------------
# Dateien
# ---------------------------------------------------------------------------

def list_directory(env: Environment, path: str = '~', show_hidden: bool = False) -> Result:
    folder = env.resolve(path, must_exist=True)
    if not folder.is_dir():
        raise ToolError(f'Kein Ordner: {folder}')
    entries = []
    for item in sorted(folder.iterdir(), key=lambda p: (not p.is_dir(), p.name.lower()))[:300]:
        if item.name.startswith('.') and not show_hidden:
            continue
        kind = 'Ordner' if item.is_dir() else f'{item.stat().st_size} Bytes'
        entries.append(f'{item.name}{"/" if item.is_dir() else ""} ({kind})')
    return Result(True, f'{folder}:\n' + ('\n'.join(entries) or '(leer)'), verified=True,
                  data={'path': str(folder), 'count': len(entries)})


def read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader  # optional, vom Installer bereitgestellt
    except ImportError:
        PdfReader = None
    if PdfReader is not None:
        reader = PdfReader(str(path))
        pages = []
        for number, page in enumerate(reader.pages[:200], 1):
            pages.append(f'[Seite {number}]\n{page.extract_text() or ""}')
            if sum(map(len, pages)) > MAX_READ:
                break
        return '\n'.join(pages)
    if shutil.which('pdftotext'):
        done = subprocess.run(['pdftotext', '-layout', str(path), '-'], capture_output=True, text=True, timeout=60)
        if done.returncode == 0:
            return done.stdout
    raise ToolError('PDF-Text kann nicht gelesen werden: pypdf fehlt (INSTALLIEREN.command erneut ausführen).')


def read_spreadsheet(path: Path) -> str:
    if path.suffix.lower() == '.csv':
        from .documents import read_csv
        return '\n'.join('\t'.join(row) for row in read_csv(path))
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ToolError('Excel-Dateien benötigen openpyxl (INSTALLIEREN.command erneut ausführen).') from exc
    book = load_workbook(str(path), read_only=True, data_only=True)
    out = []
    for sheet in book.worksheets[:10]:
        out.append(f'[Blatt {sheet.title}]')
        for row in sheet.iter_rows(max_row=500, values_only=True):
            out.append('\t'.join('' if v is None else str(v) for v in row))
    return '\n'.join(out)


def read_file(env: Environment, path: str, max_chars: int = 20_000) -> Result:
    file = env.resolve(path, must_exist=True)
    if file.is_dir():
        return list_directory(env, str(file))
    suffix = file.suffix.lower()
    if suffix == '.pdf':
        text = read_pdf(file)
    elif suffix in ('.xlsx', '.xlsm', '.csv'):
        text = read_spreadsheet(file)
    elif suffix == '.docx' and importlib.util.find_spec('docx'):
        from .documents import read_docx
        text = read_docx(file)
    elif suffix == '.pptx':
        from .documents import read_pptx
        text = read_pptx(file)
    elif suffix in ('.docx', '.doc', '.rtf', '.pages', '.odt') and env.is_mac:
        done = env.run(['textutil', '-convert', 'txt', '-stdout', str(file)], timeout=30)
        if done.returncode != 0:
            raise ToolError('Dokument konnte nicht gelesen werden.')
        text = done.stdout
    else:
        data = file.read_bytes()[:MAX_READ * 4]
        if b'\x00' in data[:4096]:
            raise ToolError('Binärdatei – kein lesbarer Text.')
        text = data.decode('utf-8', 'replace')
    limit = max(1_000, min(int(max_chars or 20_000), MAX_READ))
    cut = len(text) > limit
    return Result(True, text[:limit] + ('\n[… gekürzt]' if cut else ''), verified=True,
                  data={'path': str(file), 'chars': len(text), 'truncated': cut})


def search_files(env: Environment, query: str, folder: str = '~', limit: int = 30) -> Result:
    base = env.resolve(folder, must_exist=True)
    if not query or len(query) > 200:
        raise ToolError('Suchbegriff fehlt.')
    found: list[str] = []
    if env.is_mac and shutil.which('mdfind'):
        done = env.run(['mdfind', '-onlyin', str(base), query], timeout=20)
        found = [line for line in done.stdout.splitlines() if line][:limit]
    if not found:
        needle = query.lower()
        for root, dirs, files in os.walk(base):
            dirs[:] = [d for d in dirs if not d.startswith('.') and d not in ('Library', 'node_modules')]
            for name in dirs + files:
                if needle in name.lower():
                    found.append(str(Path(root) / name))
                    if len(found) >= limit:
                        break
            if len(found) >= limit:
                break
    return Result(True, '\n'.join(found) or 'Keine Treffer.', verified=True, data={'count': len(found)})


def create_folder(env: Environment, path: str) -> Result:
    folder = env.resolve(path, for_write=True)
    existed = folder.exists()
    folder.mkdir(parents=True, exist_ok=True)
    if not folder.is_dir():
        raise ToolError('Ordner konnte nicht angelegt werden.')
    return Result(True, f'{"Ordner existierte bereits" if existed else "Ordner angelegt"}: {folder}', verified=True,
                  data={'path': str(folder)})


def write_file(env: Environment, path: str, content: str, overwrite: bool = False) -> Result:
    file = env.resolve(path, for_write=True)
    if not isinstance(content, str):
        raise ToolError('Inhalt muss Text sein.')
    if file.exists() and not overwrite:
        raise ToolError(f'Datei existiert bereits: {file}. Zum Ersetzen overwrite=true setzen (mit Bestätigung).')
    file.parent.mkdir(parents=True, exist_ok=True)
    suffix = file.suffix.lower()
    if suffix == '.docx' and importlib.util.find_spec('docx'):
        from .documents import docx_bytes, read_docx_bytes
        data = docx_bytes(file.stem.replace('_', ' '), content)
        file.write_bytes(data)
        first = next((line.strip().lstrip('#-* ').strip() for line in content.splitlines() if line.strip()), '')
        ok = first in read_docx_bytes(file.read_bytes())
    elif suffix in ('.docx', '.rtf') and env.is_mac:
        draft = file.with_suffix('.atlas-tmp.txt')
        draft.write_text(content, encoding='utf-8')
        try:
            done = env.run(['textutil', '-convert', suffix[1:], '-output', str(file), str(draft)], timeout=30)
            if done.returncode != 0:
                raise ToolError('Dokument konnte nicht erzeugt werden.')
        finally:
            draft.unlink(missing_ok=True)
        ok = file.is_file() and file.stat().st_size > 0
    else:
        file.write_text(content, encoding='utf-8')
        ok = file.read_text(encoding='utf-8') == content
    if not ok:
        raise ToolError('Datei wurde geschrieben, die Kontrolle ist aber fehlgeschlagen.')
    return Result(True, f'Gespeichert: {file} ({file.stat().st_size} Bytes)', verified=True, data={'path': str(file)})


def normalize_rows(rows: Any, columns: Any = None, csv_text: Any = None) -> list[list]:
    """Nimmt Tabellendaten so, wie kleine Modelle sie liefern: Listen, Objekte, CSV-Text oder JSON-Text."""
    if isinstance(rows, str) and rows.strip().startswith('['):
        try:
            rows = json.loads(rows)
        except ValueError:
            pass
    if csv_text is None and isinstance(rows, str):
        csv_text, rows = rows, None
    table: list[list] = []
    if isinstance(csv_text, str) and csv_text.strip():
        sample = csv_text.strip().splitlines()[0]
        separator = ';' if sample.count(';') >= sample.count(',') else ','
        table = [row for row in csv.reader(csv_text.strip().splitlines(), delimiter=separator) if row]
    elif isinstance(rows, list):
        if rows and all(isinstance(r, dict) for r in rows):
            header = list(columns) if isinstance(columns, list) and columns else list(dict.fromkeys(k for r in rows for k in r))
            table = [header] + [[r.get(k) for k in header] for r in rows]
            columns = None
        else:
            table = [r if isinstance(r, list) else [r] for r in rows]
    if isinstance(columns, list) and columns and (not table or [str(c) for c in table[0]] != [str(c) for c in columns]):
        table = [list(columns)] + table
    return [[cell.strip() if isinstance(cell, str) else cell for cell in row] for row in table]


def create_spreadsheet(env: Environment, path: str, sheets: list | None = None, rows: Any = None,
                       columns: list | None = None, csv_text: str | None = None, overwrite: bool = False) -> Result:
    """Legt eine Excel-Tabelle an: columns + rows, CSV-Text oder mehrere Blätter (sheets)."""
    file = env.resolve(path, for_write=True)
    if file.suffix.lower() not in ('.xlsx', '.csv'):
        file = file.with_suffix('.xlsx')
    if file.exists() and not overwrite:
        raise ToolError(f'Datei existiert bereits: {file}')
    if isinstance(sheets, list) and sheets:
        sheets = [{'name': sh.get('name') or f'Tabelle{i}', 'rows': normalize_rows(sh.get('rows'), sh.get('columns'))}
                  for i, sh in enumerate(sheets, 1) if isinstance(sh, dict)]
    if not sheets:
        sheets = [{'name': 'Tabelle1', 'rows': normalize_rows(rows, columns, csv_text)}]
    if not any(sheet['rows'] for sheet in sheets):
        raise ToolError('Keine Tabellendaten. Bitte columns und rows (oder csv_text) angeben.')
    if not isinstance(sheets, list) or not all(isinstance(s, dict) for s in sheets):
        raise ToolError('sheets muss eine Liste von {name, rows} sein.')
    file.parent.mkdir(parents=True, exist_ok=True)
    total = 0
    if file.suffix.lower() == '.xlsx':
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
        except ImportError:
            file = file.with_suffix('.csv')
        else:
            book = Workbook()
            book.remove(book.active)
            for sheet in sheets[:20]:
                ws = book.create_sheet(str(sheet.get('name') or 'Tabelle')[:31])
                for index, row in enumerate(sheet.get('rows') or []):
                    values = [cell_value(v) for v in (row if isinstance(row, list) else [row])]
                    ws.append(values)
                    total += 1
                    if index == 0:
                        for cell in ws[1]:
                            cell.font = Font(bold=True)
                for column in ws.columns:
                    width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
                    ws.column_dimensions[column[0].column_letter].width = min(60, width + 2)
            book.save(str(file))
            check = read_spreadsheet(file)
            names = ', '.join(str(s.get('name') or 'Tabelle') for s in sheets)
            return Result(True, f'Excel-Tabelle gespeichert: {file} (Blätter: {names}, {total} Zeilen)',
                          verified='[Blatt' in check, data={'path': str(file), 'rows': total})
    with file.open('w', newline='', encoding='utf-8') as handle:
        writer = csv.writer(handle, delimiter=';')
        for sheet in sheets:
            for row in sheet.get('rows') or []:
                writer.writerow(row if isinstance(row, list) else [row])
                total += 1
    return Result(True, f'Als CSV gespeichert (Excel öffnet sie): {file} ({total} Zeilen)',
                  verified=file.is_file(), data={'path': str(file), 'rows': total})


def cell_value(value: Any) -> Any:
    """Formeln (=SUMME(...)) und Zahlen bleiben erhalten, alles andere wird Text."""
    if isinstance(value, (int, float)) or value is None:
        return value
    text = str(value)
    if re.fullmatch(r'-?\d+([.,]\d+)?', text.strip()):
        return float(text.replace(',', '.')) if re.search('[.,]', text) else int(text)
    return text


CATEGORIES = {
    'PDF': {'.pdf'}, 'Bilder': {'.jpg', '.jpeg', '.png', '.gif', '.heic', '.webp', '.tiff', '.bmp', '.svg'},
    'Texte': {'.txt', '.md', '.rtf'}, 'Dokumente': {'.doc', '.docx', '.pages', '.odt'},
    'Tabellen': {'.xls', '.xlsx', '.csv', '.numbers'}, 'Präsentationen': {'.ppt', '.pptx', '.key'},
    'Audio': {'.mp3', '.m4a', '.wav', '.aac', '.flac'}, 'Videos': {'.mp4', '.mov', '.m4v', '.avi', '.mkv'},
    'Archive': {'.zip', '.rar', '.7z', '.tar', '.gz', '.dmg'}, 'Code': {'.py', '.swift', '.js', '.ts', '.html', '.css', '.json'},
}


def organize_folder(env: Environment, path: str, by: str = 'typ') -> Result:
    """Sortiert die Dateien eines Ordners in Unterordner nach Dateityp (PDF, Bilder, Texte …)."""
    folder = env.resolve(path, must_exist=True, for_write=True)
    if not folder.is_dir():
        raise ToolError('Kein Ordner.')
    if folder in env.roots or folder == env.home:
        raise ToolError('Den gesamten Benutzerordner sortiere ich nicht – bitte einen konkreten Unterordner nennen.')
    moved: dict[str, list[str]] = {}
    for file in sorted(p for p in folder.iterdir() if p.is_file() and not p.name.startswith('.')):
        category = next((name for name, suffixes in CATEGORIES.items() if file.suffix.lower() in suffixes), 'Sonstiges')
        target = folder / category / file.name
        if target.exists():
            target = target.with_name(f'{file.stem}-{int(time.time())}{file.suffix}')
        target.parent.mkdir(exist_ok=True)
        shutil.move(str(file), str(target))
        moved.setdefault(category, []).append(file.name)
    loose = [p.name for p in folder.iterdir() if p.is_file() and not p.name.startswith('.')]
    if not moved:
        return Result(True, 'Keine losen Dateien zum Sortieren gefunden.', verified=True, data={'moved': {}})
    summary = '; '.join(f'{name}: {", ".join(files)}' for name, files in moved.items())
    return Result(not loose, f'Sortiert in {folder}: {summary}', verified=not loose,
                  data={'moved': moved, 'remaining': loose})


def replace_in_file(env: Environment, path: str, old: str, new: str) -> Result:
    """Ersetzt eine eindeutige Textstelle in einer Datei (für Fehlerkorrekturen im Code)."""
    file = env.resolve(path, must_exist=True, for_write=True)
    if not isinstance(old, str) or not old or not isinstance(new, str):
        raise ToolError('old (bisheriger Text) und new (neuer Text) angeben.')
    text = file.read_text(encoding='utf-8')
    count = text.count(old)
    if count == 0:
        # Häufig: Einrückung/Leerzeichen weichen ab – dann zeilenweise ohne führende Leerzeichen suchen.
        stripped = old.strip()
        lines = [line for line in text.splitlines() if line.strip() == stripped]
        if len(lines) == 1:
            old, count = lines[0], 1
            indent = lines[0][:len(lines[0]) - len(lines[0].lstrip())]
            new = indent + new.strip() if '\n' not in new.strip() else new
    if count != 1:
        raise ToolError(f'Die Textstelle kommt {count}-mal vor – sie muss genau einmal vorkommen. '
                        'Erst mit read_file den genauen Wortlaut lesen.')
    updated = text.replace(old, new, 1)
    file.write_text(updated, encoding='utf-8')
    ok = file.read_text(encoding='utf-8') == updated
    return Result(ok, f'Geändert in {file}:\n- {old.strip()[:200]}\n+ {new.strip()[:200]}', verified=ok,
                  data={'path': str(file)})


def run_tests(env: Environment, path: str, timeout: int = 180) -> Result:
    """Erkennt den Testbefehl eines Projekts und führt die Tests aus."""
    overview = inspect_project(env, path)
    commands = overview.data.get('test_commands') or []
    command = next((c for c in commands if c.startswith(('python3 -m', 'swift test', 'npm test', 'cargo test', 'go test'))), None)
    if command is None:
        root = env.resolve(path, must_exist=True)
        if any(root.rglob('test_*.py')) or any(root.rglob('*_test.py')):
            command = 'python3 -m unittest discover -s tests -t .' if (root / 'tests').is_dir() else 'python3 -m unittest discover'
    if command is None:
        raise ToolError('Kein Testbefehl erkannt. Bitte den Befehl nennen (run_command).')
    if command == 'python3 -m unittest discover -s tests':
        command += ' -t .'
    result = run_command(env, command, folder=path, timeout=timeout)
    result.data['command'] = command
    passed = result.ok
    summary = 'Alle Tests bestanden.' if passed else 'Tests fehlgeschlagen.'
    return Result(passed, f'{summary} Befehl: {command}\n' + result.output, verified=True, data=result.data)


def recent_files(env: Environment, folder: str = '~', days: int = 1, limit: int = 30) -> Result:
    """Zuletzt geänderte Dateien (z. B. „die Datei von gestern“)."""
    base = env.resolve(folder, must_exist=True)
    horizon = time.time() - max(1, min(int(days or 1), 60)) * 86400
    found = []
    for root, dirs, files in os.walk(base):
        dirs[:] = [d for d in dirs if not d.startswith('.') and d not in ('Library', 'node_modules', '__pycache__')]
        for name in files:
            if name.startswith('.'):
                continue
            item = Path(root) / name
            try:
                stamp = item.stat().st_mtime
            except OSError:
                continue
            if stamp >= horizon:
                found.append((stamp, item))
    found.sort(reverse=True)
    lines = [f'{datetime.fromtimestamp(t):%d.%m. %H:%M}  {p}' for t, p in found[:max(1, min(int(limit or 30), 100))]]
    return Result(True, '\n'.join(lines) or 'Keine Dateien in diesem Zeitraum geändert.', verified=True,
                  data={'count': len(found)})


def move_path(env: Environment, source: str, destination: str) -> Result:
    src = env.resolve(source, must_exist=True, for_write=True)
    dst = env.resolve(destination, for_write=True)
    if dst.is_dir():
        dst = dst / src.name
    if dst.exists():
        raise ToolError(f'Ziel existiert bereits: {dst}')
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.move(str(src), str(dst))
    return Result(dst.exists() and not src.exists(), f'Verschoben nach {dst}', verified=dst.exists() and not src.exists(),
                  data={'path': str(dst)})


def delete_path(env: Environment, path: str) -> Result:
    """Löschen heißt: in den Papierkorb legen – wiederherstellbar."""
    target = env.resolve(path, must_exist=True, for_write=True)
    if target in env.roots or target == env.home:
        raise ToolError('Freigegebene Hauptordner werden nie gelöscht.')
    if env.is_mac:
        script = f'tell application "Finder" to delete (POSIX file {json.dumps(str(target))} as alias)'
        done = env.run(['osascript', '-e', script], timeout=20)
        ok = done.returncode == 0 and not target.exists()
        if not ok:
            raise ToolError('Finder konnte die Datei nicht in den Papierkorb legen: ' + (done.stderr or '')[:200])
    else:
        env.trash.mkdir(exist_ok=True)
        destination = env.trash / f'{target.name}-{int(time.time())}'
        shutil.move(str(target), str(destination))
        ok = not target.exists()
    return Result(ok, f'In den Papierkorb gelegt: {target}', verified=ok)


# ---------------------------------------------------------------------------
# Programme, Finder, Browser, Web
# ---------------------------------------------------------------------------

APP_ALIASES = {
    'rechner': 'Calculator', 'taschenrechner': 'Calculator', 'notizen': 'Notes', 'kalender': 'Calendar',
    'erinnerungen': 'Reminders', 'mail': 'Mail', 'nachrichten': 'Messages', 'fotos': 'Photos',
    'musik': 'Music', 'karten': 'Maps', 'kontakte': 'Contacts', 'systemeinstellungen': 'System Settings',
    'einstellungen': 'System Settings', 'terminal': 'Terminal', 'vorschau': 'Preview', 'textedit': 'TextEdit',
    'excel': 'Microsoft Excel', 'word': 'Microsoft Word', 'chrome': 'Google Chrome', 'finder': 'Finder',
    'safari': 'Safari', 'spotify': 'Spotify', 'xcode': 'Xcode', 'numbers': 'Numbers', 'pages': 'Pages',
}
APP_FOLDERS = [Path('/Applications'), Path('/System/Applications'), Path('/System/Applications/Utilities'),
               Path.home() / 'Applications']


def find_app(env: Environment, name: str) -> str | None:
    wanted = APP_ALIASES.get(name.strip().lower(), name.strip())
    if wanted.lower() == 'finder':
        return 'Finder'
    for folder in APP_FOLDERS:
        if not folder.is_dir():
            continue
        for app in folder.glob('*.app'):
            if app.stem.lower() == wanted.lower():
                return app.stem
    for folder in APP_FOLDERS:
        if folder.is_dir():
            for app in folder.glob('*.app'):
                if wanted.lower() in app.stem.lower():
                    return app.stem
    return wanted if not env.is_mac else None


def open_app(env: Environment, name: str) -> Result:
    if not name or len(name) > 80 or '/' in name:
        raise ToolError('Ungültiger Programmname.')
    app = find_app(env, name)
    if not app:
        raise ToolError(f'Programm „{name}“ ist nicht installiert.')
    done = env.run(['open', '-a', app], timeout=15)
    if done.returncode != 0:
        raise ToolError(f'„{app}“ ließ sich nicht öffnen: {(done.stderr or "").strip()[:200]}')
    running = False
    for _ in range(10):
        check = env.run(['pgrep', '-xq', app], timeout=5)
        if check.returncode == 0:
            running = True
            break
        time.sleep(0.3)
    return Result(True, f'{app} geöffnet' + ('' if running else ' (Start angestoßen, läuft noch nicht sichtbar)'),
                  verified=running, data={'app': app})


def open_path(env: Environment, path: str, reveal: bool = False) -> Result:
    target = env.resolve(path, must_exist=True)
    args = ['open', '-R', str(target)] if reveal else ['open', str(target)]
    done = env.run(args, timeout=15)
    if done.returncode != 0:
        raise ToolError('Öffnen fehlgeschlagen: ' + (done.stderr or '')[:200])
    return Result(True, f'Geöffnet: {target}', verified=True, data={'path': str(target)})


def safe_url(url: str) -> str:
    parsed = urllib.parse.urlparse(url.strip())
    if parsed.scheme not in ('http', 'https') or not parsed.netloc:
        raise ToolError('Nur http- und https-Adressen sind erlaubt.')
    host = parsed.hostname or ''
    if host in ('localhost',) or host.startswith('127.') or host.endswith('.local') or host.startswith('192.168.'):
        raise ToolError('Adressen im lokalen Netz werden nicht geöffnet.')
    return parsed.geturl()


def open_url(env: Environment, url: str, browser: str = 'Safari') -> Result:
    target = safe_url(url)
    app = find_app(env, browser or 'Safari') or 'Safari'
    done = env.run(['open', '-a', app, target], timeout=15)
    if done.returncode != 0:
        raise ToolError('Browser ließ sich nicht öffnen.')
    return Result(True, f'{app} zeigt {target}', verified=True, data={'url': target})


def search_web(env: Environment, query: str, open_in_browser: bool = False, max_results: int = 6) -> Result:
    if not query or len(query) > 300:
        raise ToolError('Suchbegriff fehlt.')
    search_url = 'https://duckduckgo.com/?q=' + urllib.parse.quote(query)
    if open_in_browser:
        open_url(env, search_url)
    raw = env.fetch('https://html.duckduckgo.com/html/?q=' + urllib.parse.quote(query), 15)
    results = []
    for match in re.finditer(r'(?s)<a[^>]+class="result__a"[^>]+href="([^"]+)"[^>]*>(.*?)</a>(.*?)'
                             r'(?:<a[^>]+class="result__snippet"[^>]*>(.*?)</a>|</div>)', raw):
        link = html.unescape(match[1])
        if 'uddg=' in link:
            link = urllib.parse.unquote(urllib.parse.parse_qs(urllib.parse.urlparse(link).query).get('uddg', [link])[0])
        title = html_to_text(match[2])
        snippet = html_to_text(match[4] or '')
        results.append({'title': title, 'url': link, 'snippet': snippet})
        if len(results) >= max(1, min(int(max_results or 6), 10)):
            break
    if not results:
        return Result(False, 'Keine Suchergebnisse erhalten (Internet oder Suchdienst nicht erreichbar).',
                      data={'query': query})
    text = '\n'.join(f'{i}. {r["title"]} – {r["url"]}\n   {r["snippet"]}' for i, r in enumerate(results, 1))
    return Result(True, text + ('\n(Auch im Browser geöffnet.)' if open_in_browser else ''), verified=True,
                  data={'results': results})


def fetch_webpage(env: Environment, url: str, max_chars: int = 12_000) -> Result:
    text = html_to_text(env.fetch(safe_url(url), 20))
    limit = max(1_000, min(int(max_chars or 12_000), 40_000))
    return Result(bool(text), text[:limit] or 'Seite ohne lesbaren Text.', verified=bool(text),
                  data={'url': url, 'chars': len(text)})


# ---------------------------------------------------------------------------
# System, Terminal, Bildschirm, Rechnen
# ---------------------------------------------------------------------------

def system_info(env: Environment) -> Result:
    info = {'Gerät': platform.machine(), 'System': platform.platform(), 'Zeit': datetime.now().strftime('%d.%m.%Y %H:%M')}
    usage = shutil.disk_usage(env.home)
    info['Speicher frei'] = f'{usage.free / 1e9:.1f} GB von {usage.total / 1e9:.1f} GB'
    if env.is_mac:
        for key, args in [('macOS', ['sw_vers', '-productVersion']), ('Akku', ['pmset', '-g', 'batt']),
                          ('Arbeitsspeicher', ['sysctl', '-n', 'hw.memsize']), ('Chip', ['sysctl', '-n', 'machdep.cpu.brand_string'])]:
            try:
                done = env.run(args, timeout=5)
            except (OSError, subprocess.TimeoutExpired):
                continue
            value = done.stdout.strip()
            if key == 'Arbeitsspeicher' and value.isdigit():
                value = f'{int(value) / 2**30:.0f} GB'
            if key == 'Akku':
                found = re.search(r'(\d+%;[^;]*)', value)
                value = found[1] if found else value.splitlines()[-1] if value else ''
            if value:
                info[key] = value
    try:
        load = os.getloadavg()
        info['Auslastung'] = f'{load[0]:.2f} (1 min)'
    except OSError:
        pass
    return Result(True, '\n'.join(f'{k}: {v}' for k, v in info.items()), verified=True, data=info)


# Terminal: KEINE freie Shell. Nur Programme aus dieser Liste, ohne Shell ausgeführt (keine Pipes,
# Umleitungen, Variablen oder Befehlsketten), Pfade nur innerhalb der freigegebenen Ordner.
READ_ONLY_COMMANDS = {
    'ls', 'pwd', 'cat', 'head', 'tail', 'wc', 'grep', 'rg', 'file', 'stat', 'du', 'df', 'date', 'uname', 'sw_vers',
    'whoami', 'tree', 'diff', 'sort', 'uniq', 'md5', 'shasum', 'which', 'find',
}
GIT_READ = {'status', 'log', 'diff', 'show', 'rev-parse', 'ls-files', 'blame', 'branch', 'remote'}
GIT_CONFIRM = {'add', 'commit', 'checkout', 'switch', 'stash', 'restore', 'pull', 'push', 'fetch', 'tag', 'init', 'mv'}
FIND_FORBIDDEN = {'-delete', '-exec', '-execdir', '-ok', '-okdir', '-fprint', '-fprintf', '-fls'}
SHELL_META = re.compile(r'[;&|`$<>\n\\]')


def command_risk(args: dict, env: 'Environment | None' = None) -> str:
    """SAFE (lesend), CONFIRM (bauen, testen, Git ändern) oder BLOCKED (alles andere)."""
    command = str(args.get('command', ''))
    if not command.strip() or len(command) > 2000 or SHELL_META.search(command):
        return BLOCKED
    try:
        parts = shlex.split(command)
    except ValueError:
        return BLOCKED
    program, rest = parts[0], parts[1:]
    if '/' in program:
        return BLOCKED
    if env is not None and not paths_allowed(env, args.get('folder') or '~', rest):
        return BLOCKED
    if program in READ_ONLY_COMMANDS:
        if program == 'find' and any(p in FIND_FORBIDDEN for p in rest):
            return BLOCKED
        return SAFE
    sub = rest[0] if rest else ''
    if program == 'git':
        if sub == 'branch' and any(p in ('-d', '-D', '-m', '-M', '--delete') for p in rest):
            return CONFIRM
        if sub == 'remote' and len(rest) > 1 and rest[1] not in ('-v', 'show', 'get-url'):
            return BLOCKED
        if sub in GIT_READ:
            return SAFE
        if sub in GIT_CONFIRM and not any(p in ('--force', '-f', '--hard', '--force-with-lease') for p in rest):
            return CONFIRM
        return BLOCKED
    if program in ('python3', 'python'):
        if rest in (['--version'], ['-V']):
            return SAFE
        if len(rest) >= 2 and rest[0] == '-m' and rest[1] in ('pytest', 'unittest', 'py_compile', 'compileall'):
            return CONFIRM
        if rest and rest[0].endswith('.py') and not rest[0].startswith('-'):
            return CONFIRM          # ein Skript des Projekts ausführen
        return BLOCKED              # kein python -c, kein pip, keine Module nach Wahl
    rules = {
        'swift': ({'--version'}, {'build', 'test', 'run'}),
        'xcodebuild': ({'-version', '-list'}, {'build', 'test', '-scheme', '-project', '-workspace', 'clean'}),
        'npm': ({'ls', 'view', '--version', 'outdated'}, {'test', 'run', 'install', 'ci'}),
        'brew': ({'list', 'info', '--version', 'outdated'}, set()),
        'make': (set(), {sub} if sub and not sub.startswith('-') else {''}),
        'cargo': ({'--version'}, {'build', 'test', 'check', 'run'}),
        'go': ({'version', 'env'}, {'build', 'test', 'vet', 'run'}),
        'swiftc': ({'--version'}, set()),
    }
    if program in rules:
        safe, confirm = rules[program]
        if sub in safe:
            return SAFE
        if sub in confirm or (program == 'xcodebuild' and any(p in ('build', 'test') for p in rest)):
            return CONFIRM
    return BLOCKED


def paths_allowed(env: 'Environment', folder: str, arguments: list[str]) -> bool:
    """Alle Pfad-Argumente müssen innerhalb der freigegebenen Ordner liegen."""
    try:
        cwd = env.resolve(folder, must_exist=True)
    except ToolError:
        return False
    for argument in arguments:
        value = argument.split('=', 1)[1] if argument.startswith('--') and '=' in argument else argument
        if value.startswith('-') or not value:
            continue
        if value.startswith(('/', '~')) or '..' in Path(value).parts:
            try:
                candidate = value if value.startswith(('/', '~')) else str(cwd / value)
                env.resolve(candidate)
            except ToolError:
                return False
    return True


def run_command(env: Environment, command: str, folder: str = '~', timeout: int = 60) -> Result:
    """Führt einen Befehl aus der erlaubten Liste ohne Shell aus (Argumente als Liste)."""
    if command_risk({'command': command, 'folder': folder}, env) == BLOCKED:
        raise ToolError('Dieser Befehl ist nicht erlaubt. Erlaubt sind lesende Befehle (ls, cat, grep, git status …) '
                        'sowie Bauen/Testen (python3 -m pytest/unittest, swift, xcodebuild, npm, cargo, go, make, git) – '
                        'ohne Pipes, Umleitungen oder sudo.')
    cwd = env.resolve(folder, must_exist=True)
    if not cwd.is_dir():
        raise ToolError('Arbeitsordner ist kein Ordner.')
    limit = max(5, min(int(timeout or 60), 600))
    started = time.monotonic()
    try:
        # Eigener, leerer Bytecode-Ordner je Lauf: Sonst testet Python nach einer Korrektur in derselben
        # Sekunde womöglich den alten Code aus __pycache__.
        with tempfile.TemporaryDirectory(prefix='atlas-run-') as cache:
            done = env.run(shlex.split(command), timeout=limit, cwd=cwd,
                           extra_env={'PYTHONDONTWRITEBYTECODE': '1', 'PYTHONPYCACHEPREFIX': cache, 'NO_COLOR': '1'})
    except subprocess.TimeoutExpired:
        raise ToolError(f'Zeitlimit von {limit} s überschritten.')
    except FileNotFoundError:
        raise ToolError('Programm ist auf diesem Mac nicht installiert.')
    output = (done.stdout or '') + (('\n[stderr]\n' + done.stderr) if done.stderr else '')
    seconds = round(time.monotonic() - started, 1)
    text = f'Exit-Code {done.returncode} nach {seconds} s\n' + output[-MAX_OUTPUT:]
    return Result(done.returncode == 0, text, verified=True, data={'exit_code': done.returncode})


def take_screenshot(env: Environment, path: str = '') -> Result:
    if not env.is_mac:
        raise ToolError('Bildschirmfotos gibt es nur auf dem Mac.')
    target = env.resolve(path or f'~/Desktop/Atlas-Bildschirm-{datetime.now():%Y%m%d-%H%M%S}.png', for_write=True)
    target.parent.mkdir(parents=True, exist_ok=True)
    done = env.run(['screencapture', '-x', str(target)], timeout=20)
    ok = done.returncode == 0 and target.is_file() and target.stat().st_size > 0
    if not ok:
        raise ToolError('Bildschirmfoto fehlgeschlagen. macOS verlangt dafür die Freigabe „Bildschirmaufnahme“ für Python '
                        '(Systemeinstellungen → Datenschutz & Sicherheit).')
    return Result(True, f'Bildschirmfoto gespeichert: {target}', verified=True, data={'path': str(target)})


VISION_HINT = ('Für Bildanalyse fehlt ein lokales Bildmodell. Nachrüsten im Terminal: "ollama pull gemma3:4b" '
               '(ca. 3 GB) – danach versteht ATLAS Bildschirm und Bilder.')


def ask_vision(env: Environment, images: list[bytes], question: str) -> str:
    if env.vision is None:
        raise ToolError(VISION_HINT)
    answer = env.vision(images, question or 'Beschreibe kurz und sachlich, was zu sehen ist. Nenne lesbare Texte.')
    if not answer:
        raise ToolError(VISION_HINT)
    return answer.strip()


def analyze_screen(env: Environment, question: str = '') -> Result:
    """Bildschirmfoto (nur temporär, wird sofort gelöscht) und Analyse mit dem lokalen Bildmodell."""
    if not env.is_mac:
        raise ToolError('Bildschirmanalyse gibt es nur auf dem Mac.')
    if env.vision is None:
        raise ToolError(VISION_HINT)
    import tempfile
    with tempfile.TemporaryDirectory() as folder:
        shot = Path(folder) / 'screen.png'
        done = env.run(['screencapture', '-x', '-t', 'png', str(shot)], timeout=20)
        if done.returncode != 0 or not shot.is_file() or shot.stat().st_size == 0:
            raise ToolError('Bildschirmfoto fehlgeschlagen. macOS verlangt dafür die Freigabe „Bildschirmaufnahme“ für Python '
                            '(Systemeinstellungen → Datenschutz & Sicherheit).')
        answer = ask_vision(env, [shot.read_bytes()], question)
    return Result(True, answer, verified=True, data={'quelle': 'Bildschirm (Foto bereits gelöscht)'})


def analyze_image(env: Environment, path: str, question: str = '') -> Result:
    file = env.resolve(path, must_exist=True)
    if file.suffix.lower() not in ('.png', '.jpg', '.jpeg', '.heic', '.gif', '.webp', '.bmp', '.tiff'):
        raise ToolError('Keine Bilddatei.')
    if file.stat().st_size > 20_000_000:
        raise ToolError('Bild ist zu groß (über 20 MB).')
    return Result(True, ask_vision(env, [file.read_bytes()], question), verified=True, data={'path': str(file)})


def safari_current_page(env: Environment, max_chars: int = 8_000) -> Result:
    """Liest Adresse, Titel und Text der vorne geöffneten Safari-Seite (macOS fragt einmal nach Erlaubnis)."""
    if not env.is_mac:
        raise ToolError('Safari gibt es nur auf dem Mac.')
    script = 'tell application "Safari" to return (URL of front document) & linefeed & (name of front document)'
    done = env.run(['osascript', '-e', script], timeout=10)
    if done.returncode != 0 or not done.stdout.strip():
        raise ToolError('Keine Safari-Seite geöffnet oder Steuerung von Safari nicht erlaubt '
                        '(Systemeinstellungen → Datenschutz & Sicherheit → Automation).')
    url, _, title = done.stdout.strip().partition('\n')
    text = ''
    try:
        text = html_to_text(env.fetch(safe_url(url), 20))
    except (ToolError, OSError, ValueError):
        text = '(Seiteninhalt konnte nicht geladen werden – evtl. Anmeldung nötig oder lokale Seite.)'
    limit = max(1_000, min(int(max_chars or 8_000), 30_000))
    return Result(True, f'{title.strip()}\n{url.strip()}\n\n{text[:limit]}', verified=True,
                  data={'url': url.strip(), 'title': title.strip()})


PROJECT_MARKERS = [
    ('Package.swift', 'Swift-Paket', 'swift test'),
    ('*.xcodeproj', 'Xcode-Projekt', 'xcodebuild -list'),
    ('pyproject.toml', 'Python', 'python3 -m pytest'),
    ('setup.py', 'Python', 'python3 -m pytest'),
    ('requirements.txt', 'Python', 'python3 -m pytest'),
    ('package.json', 'JavaScript/Node', 'npm test'),
    ('Cargo.toml', 'Rust', 'cargo test'),
    ('go.mod', 'Go', 'go test ./...'),
]


def inspect_project(env: Environment, path: str) -> Result:
    """Überblick über ein Softwareprojekt: Art, Struktur, README, Tests, Git-Status."""
    root = env.resolve(path, must_exist=True)
    if not root.is_dir():
        raise ToolError('Kein Projektordner.')
    kinds, commands = [], []
    for pattern, kind, command in PROJECT_MARKERS:
        if list(root.glob(pattern)):
            kinds.append(kind)
            commands.append(command)
    if (root / 'package.json').is_file():
        try:
            scripts = json.loads((root / 'package.json').read_text(encoding='utf-8')).get('scripts', {})
            if 'test' not in scripts:
                commands = [c for c in commands if c != 'npm test']
        except ValueError:
            pass
    if 'Python' in kinds and not any(root.rglob('conftest.py')) and (root / 'tests').is_dir():
        commands = [c.replace('python3 -m pytest', 'python3 -m unittest discover -s tests') for c in commands]
    tree = []
    for item in sorted(root.rglob('*')):
        relative = item.relative_to(root)
        if any(part.startswith('.') or part in ('node_modules', '__pycache__', 'build', 'DerivedData', 'venv', '.venv')
               for part in relative.parts):
            continue
        if len(relative.parts) <= 2:
            tree.append(('  ' * (len(relative.parts) - 1)) + relative.name + ('/' if item.is_dir() else ''))
        if len(tree) >= 80:
            tree.append('…')
            break
    readme = next((p for p in root.iterdir() if p.name.lower().startswith('readme')), None)
    summary = [f'Projekt: {root}', 'Art: ' + (', '.join(dict.fromkeys(kinds)) or 'unbekannt'),
               'Vorgeschlagene Tests: ' + (', '.join(dict.fromkeys(commands)) or 'keine erkannt')]
    if (root / '.git').exists():
        done = env.run(['git', 'status', '--short', '--branch'], timeout=10, cwd=root)
        summary.append('Git:\n' + (done.stdout.strip()[:1500] or '(sauber)'))
    summary.append('Struktur:\n' + '\n'.join(tree))
    if readme:
        summary.append('README (Anfang):\n' + readme.read_text(encoding='utf-8', errors='replace')[:1500])
    return Result(True, '\n\n'.join(summary), verified=True,
                  data={'kinds': list(dict.fromkeys(kinds)), 'test_commands': list(dict.fromkeys(commands))})


def calculate_tool(env: Environment, expression: str) -> Result:
    answer = calculate(normalize_expression(str(expression)))
    if answer is None:
        raise ToolError('Ausdruck nicht berechenbar (erlaubt: Zahlen, + − × ÷, Klammern, Prozent).')
    return Result(True, answer, verified=True)


# ---------------------------------------------------------------------------
# Registry
# ---------------------------------------------------------------------------

def _describe_path(verb: str) -> Callable[[dict], str]:
    return lambda a: f'{verb}: {a.get("path") or a.get("source", "")}'


def build_tools(env: Environment) -> dict[str, Tool]:
    from .research import research
    from .documents import (convert_document, create_document, create_pdf, create_presentation, edit_document,
                            edit_spreadsheet)
    def bind(func: Callable[..., Result]) -> Callable[..., Result]:
        return lambda **kwargs: func(env, **kwargs)

    path = string('Pfad, z. B. "~/Documents/Bericht.txt" oder "Schreibtisch/Neu".')
    tools = [
        Tool('list_directory', 'Listet den Inhalt eines Ordners auf.', obj({'path': path}), bind(list_directory),
             category='files', label='Liest den Ordner'),
        Tool('read_file', 'Liest eine Datei: Text, Code, PDF, Word, Excel/CSV. Für PDF-Zusammenfassungen.',
             obj({'path': path, 'max_chars': {'type': 'integer', 'description': 'Höchstens so viele Zeichen.'}}, ['path']),
             bind(read_file), category='files', label='Liest die Datei'),
        Tool('search_files', 'Sucht Dateien und Ordner nach Name oder Inhalt (Spotlight).',
             obj({'query': string('Suchbegriff.'), 'folder': path}, ['query']), bind(search_files),
             category='files', label='Sucht Dateien'),
        Tool('create_folder', 'Legt einen neuen Ordner an.', obj({'path': path}, ['path']), bind(create_folder),
             category='files', label='Legt den Ordner an'),
        Tool('write_file', 'Speichert Text in eine NEUE Datei (.txt, .md, .py, .docx, .rtf …). Überschreiben nur mit overwrite=true.',
             obj({'path': path, 'content': string('Vollständiger Inhalt.'),
                  'overwrite': {'type': 'boolean', 'description': 'Vorhandene Datei ersetzen (wird bestätigt).'}},
                 ['path', 'content']),
             bind(write_file), category='files', label='Speichert die Datei',
             risk_for=lambda a: CONFIRM if a.get('overwrite') else SAFE, describe=_describe_path('Datei ersetzen')),
        Tool('create_pdf', 'Erstellt eine PDF mit Titel und vollständigem Text. Nutzt automatische Seitenumbrüche. Neuer Dateiname erforderlich.',
             obj({'path': path, 'title': string('Dokumenttitel.'), 'content': string('Vollständiger Text mit Absätzen.')},
                 ['path', 'title', 'content']), bind(create_pdf), category='documents', label='Erstellt die PDF'),
        Tool('research', 'Recherchiert im Internet: sucht, liest mehrere Seiten und liefert nummerierte Quellen mit Adresse '
                         'und Abrufdatum. Für aktuelle Informationen und Fragen, bei denen Quellen wichtig sind.',
             obj({'question': string('Suchanfrage bzw. Frage.'), 'max_sources': {'type': 'integer', 'description': '1–6, Standard 4.'}},
                 ['question']), bind(research), category='web', label='Recherchiert'),
        Tool('create_document', 'Erstellt ein echtes Word-Dokument (.docx) – für Briefe, Angebote, Berichte, Konzepte, '
                                'E-Mail-Entwürfe. content: vollständiger Text; "# " Überschrift, "## " Unterüberschrift, "- " Aufzählung.',
             obj({'path': path, 'title': string('Dokumenttitel.'), 'content': string('Vollständiger Text.')},
                 ['path', 'title', 'content']), bind(create_document), category='documents', label='Erstellt das Word-Dokument'),
        Tool('edit_document', 'Ändert ein vorhandenes Dokument (.docx, .txt, .md): Textstellen ersetzen und/oder Text anhängen. '
                              'Vorher wird automatisch eine Sicherung angelegt.',
             obj({'path': path, 'replacements': {'type': 'array', 'items': obj({'old': string('Genauer alter Text.'),
                  'new': string('Neuer Text.')}, ['old', 'new'])}, 'append': string('Text, der ans Ende kommt (optional).')},
                 ['path']), bind(edit_document), risk=CONFIRM, category='documents', label='Ändert das Dokument',
             describe=_describe_path('Dokument ändern (mit Sicherung)')),
        Tool('edit_spreadsheet', 'Ändert eine vorhandene Excel- oder CSV-Tabelle: Zellen setzen (cells, z. B. {"B2": 120, '
                                 '"B5": "=SUM(B2:B4)"}) und/oder Zeilen anhängen (append_rows). Mit Sicherung.',
             obj({'path': path, 'cells': {'type': 'object', 'description': 'Zellbezug → Wert oder Formel.'},
                  'append_rows': {'type': 'array', 'items': {'type': 'array'}}, 'sheet': string('Blattname (optional).')},
                 ['path']), bind(edit_spreadsheet), risk=CONFIRM, category='documents', label='Ändert die Tabelle',
             describe=_describe_path('Tabelle ändern (mit Sicherung)')),
        Tool('convert_document', 'Wandelt eine Datei um, ohne das Original zu ändern: .docx/.txt/.md/.csv/.xlsx/.pptx → PDF, '
                                 'oder .txt/.md → Word (to="docx"). Für „Speichere es als PDF“.',
             obj({'path': path, 'to': {'type': 'string', 'enum': ['pdf', 'docx']},
                  'target': string('Zielpfad (optional, sonst gleicher Name mit neuer Endung).')}, ['path']),
             bind(convert_document), category='documents', label='Wandelt die Datei um'),
        Tool('create_presentation', 'Erstellt eine PowerPoint-Datei (.pptx), 1 bis 40 Folien, höchstens 6 Stichpunkte pro Folie. Neuer Dateiname erforderlich.',
             obj({'path': path, 'title': string('Gesamttitel.'), 'slides': {'type': 'array', 'items':
                 obj({'title': string('Folientitel, höchstens 100 Zeichen.'), 'bullets': {'type': 'array',
                      'items': {'type': 'string'}, 'description': '1 bis 6 Stichpunkte, je höchstens 180 Zeichen.'}}, ['title', 'bullets'])}},
                 ['path', 'title', 'slides']), bind(create_presentation), category='documents', label='Erstellt die Präsentation'),
        Tool('create_spreadsheet', 'Erstellt eine Excel-Tabelle (.xlsx) mit Daten. Einfachster Weg: columns (Überschriften) '
                                   'und rows (Zeilen), z. B. columns=["Posten","Betrag"], rows=[["Miete",1200],["Strom",90]]. '
                                   'Formeln wie "=SUM(B2:B3)" sind erlaubt.',
             obj({'path': path, 'columns': {'type': 'array', 'items': {'type': 'string'}, 'description': 'Spaltenüberschriften.'},
                  'rows': {'type': 'array', 'items': {'type': 'array'}, 'description': 'Datenzeilen.'},
                  'csv_text': string('Alternativ: Tabelle als CSV-Text mit ; getrennt, erste Zeile = Überschriften.'),
                  'sheets': {'type': 'array', 'items': {'type': 'object'},
                             'description': 'Nur für mehrere Blätter: [{"name": "...", "columns": [...], "rows": [...]}]'}},
                 ['path']),
             bind(create_spreadsheet), category='documents', label='Erstellt die Tabelle'),
        Tool('organize_folder', 'Sortiert alle Dateien eines Ordners automatisch in Unterordner nach Dateityp '
                                '(PDF, Bilder, Texte, Tabellen …). Ein Aufruf erledigt alles.',
             obj({'path': path}, ['path']), bind(organize_folder), risk=CONFIRM, category='files',
             label='Sortiert den Ordner', describe=_describe_path('Dateien nach Typ in Unterordner sortieren')),
        Tool('replace_in_file', 'Ändert eine Datei gezielt: ersetzt eine Textstelle (old) durch neuen Text (new). '
                                'Ideal für Fehlerkorrekturen in Code. Vorher mit read_file den genauen Wortlaut lesen.',
             obj({'path': path, 'old': string('Bisheriger Text, genau wie in der Datei.'), 'new': string('Neuer Text.')},
                 ['path', 'old', 'new']), bind(replace_in_file), risk=CONFIRM, category='terminal', label='Ändert die Datei',
             describe=lambda a: f'Datei ändern: {a.get("path")}\n- {str(a.get("old", ""))[:120]}\n+ {str(a.get("new", ""))[:120]}'),
        Tool('run_tests', 'Erkennt den Testbefehl eines Projekts und führt die Tests aus (Python, Swift, Node, Rust, Go).',
             obj({'path': path}, ['path']), bind(run_tests), risk=CONFIRM, category='terminal', label='Führt die Tests aus',
             describe=lambda a: f'Tests im Projekt ausführen: {a.get("path")}'),
        Tool('recent_files', 'Listet zuletzt geänderte Dateien, z. B. „die Datei von gestern“ (days=1) oder der letzten Woche (days=7).',
             obj({'folder': path, 'days': {'type': 'integer'}}), bind(recent_files), category='files',
             label='Sucht zuletzt geänderte Dateien'),
        Tool('move_path', 'Verschiebt oder benennt eine Datei/einen Ordner um.',
             obj({'source': path, 'destination': path}, ['source', 'destination']), bind(move_path),
             risk=CONFIRM, category='files', label='Verschiebt',
             describe=lambda a: f'Verschieben: {a.get("source")} → {a.get("destination")}'),
        Tool('delete_path', 'Legt eine Datei oder einen Ordner in den Papierkorb.', obj({'path': path}, ['path']),
             bind(delete_path), risk=CONFIRM, category='files', label='Löscht', describe=_describe_path('In den Papierkorb legen')),
        Tool('open_app', 'Öffnet ein installiertes Programm (z. B. Safari, Notizen, Kalender, Excel).',
             obj({'name': string('Programmname.')}, ['name']), bind(open_app), category='apps', label='Öffnet das Programm'),
        Tool('open_path', 'Öffnet eine Datei oder einen Ordner im passenden Programm bzw. Finder (reveal=true zeigt sie im Finder).',
             obj({'path': path, 'reveal': {'type': 'boolean'}}, ['path']), bind(open_path), category='apps',
             label='Öffnet'),
        Tool('open_url', 'Öffnet eine Webseite im Browser.', obj({'url': string('http(s)-Adresse.'),
                                                                  'browser': string('Standard: Safari.')}, ['url']),
             bind(open_url), category='web', label='Öffnet die Webseite'),
        Tool('search_web', 'Sucht im Internet nach aktuellen Informationen und liefert Titel, Links und Auszüge. '
                           'open_in_browser=true zeigt die Suche zusätzlich in Safari.',
             obj({'query': string('Suchanfrage.'), 'open_in_browser': {'type': 'boolean'}}, ['query']),
             bind(search_web), category='web', label='Recherchiert im Internet'),
        Tool('fetch_webpage', 'Lädt den Text einer Webseite, um sie zu lesen oder zusammenzufassen.',
             obj({'url': string('http(s)-Adresse.')}, ['url']), bind(fetch_webpage), category='web', label='Liest die Webseite'),
        Tool('system_info', 'Zeigt Systeminformationen: macOS-Version, Akku, freier Speicher, Arbeitsspeicher, Auslastung.',
             obj({}), bind(system_info), category='system', label='Prüft das System'),
        Tool('run_command', 'Führt einen erlaubten Terminalbefehl in einem Projektordner aus: lesend (ls, cat, grep, find, '
                            'git status/diff/log) ohne Rückfrage; Bauen/Testen (python3 -m pytest|unittest, swift build|test, '
                            'xcodebuild, npm test, cargo, go, make, git commit) mit Bestätigung. Keine Pipes, keine Shell.',
             obj({'command': string('Befehl, z. B. "python3 -m unittest".'), 'folder': path,
                  'timeout': {'type': 'integer', 'description': 'Sekunden, Standard 60.'}}, ['command']),
             bind(run_command), category='terminal', label='Führt den Befehl aus',
             risk_for=lambda a: command_risk(a, env),
             describe=lambda a: f'Terminal ({a.get("folder", "~")}): {a.get("command")}'),
        Tool('take_screenshot', 'Macht ein Bildschirmfoto und speichert es.', obj({'path': path}), bind(take_screenshot),
             category='screen', label='Macht ein Bildschirmfoto'),
        Tool('analyze_screen', 'Schaut auf den Bildschirm (lokales Bildmodell) und beantwortet eine Frage dazu, '
                               'z. B. „Was steht in der Fehlermeldung?“. Das Foto wird sofort wieder gelöscht.',
             obj({'question': string('Was soll auf dem Bildschirm geprüft werden?')}), bind(analyze_screen),
             category='screen', label='Schaut auf den Bildschirm'),
        Tool('analyze_image', 'Analysiert eine Bilddatei (Foto, Screenshot, Scan) mit dem lokalen Bildmodell.',
             obj({'path': path, 'question': string('Frage zum Bild.')}, ['path']), bind(analyze_image),
             category='screen', label='Analysiert das Bild'),
        Tool('safari_current_page', 'Liest die gerade in Safari geöffnete Seite (Adresse, Titel, Text).',
             obj({}), bind(safari_current_page), category='web', label='Liest die Safari-Seite'),
        Tool('inspect_project', 'Verschafft einen Überblick über ein Softwareprojekt: Art, Struktur, README, '
                                'passende Testbefehle, Git-Status. Erster Schritt bei Fehlersuche.',
             obj({'path': path}, ['path']), bind(inspect_project), category='terminal', label='Untersucht das Projekt'),
        Tool('calculate', 'Rechnet exakt: + − × ÷, Klammern, Prozent ("15% von 80").',
             obj({'expression': string('Rechenausdruck.')}, ['expression']), bind(calculate_tool),
             category='core', label='Rechnet'),
    ]
    return {tool.name: tool for tool in tools}


# Schlüsselwörter, mit denen der Router passende Werkzeuggruppen auswählt (kleinere Modelle
# arbeiten zuverlässiger, wenn sie nur die relevanten Werkzeuge sehen – und schneller).
CATEGORY_HINTS = {
    'files': r'datei|ordner|folder|file|speicher|dokument|pdf|lies|lese|text|schreib|notiz|liste|lösch|verschieb|umbenenn|download|schreibtisch|desktop|projekt|code|sortier|aufräum|gestern|zuletzt',
    'documents': r'pdf|word|docx|dokument|brief|angebot|bericht|konzept|entwurf|umwandel|konvertier|als pdf|powerpoint|präsentation|presentation|folien|pptx|excel|tabelle|spreadsheet|xlsx|csv|kalkulation|budget|finanz|übersicht',
    'apps': r'öffne|starte|open|programm|app\b|finder|zeig',
    'web': r'such|recherch|internet|web|google|safari|browser|neueste|aktuell|news|nachrichten|website|seite|link|http',
    'system': r'system|akku|batterie|speicherplatz|festplatte|arbeitsspeicher|ram|cpu|mac\b|version',
    'terminal': r'terminal|befehl|kommando|fehler|bug|test|build|kompil|git|projekt|code|skript|behebe|repariere',
    'screen': r'bildschirm|screenshot|bildschirmfoto|screen|bild\b|foto|siehst du|schau|fehlermeldung',
    'learn': r'\blern|bring dir|beibring|fähigkeit|informier dich|eigne dir',
    'core': r'rechne|berechne|prozent|summe|\d\s*[+*/×÷-]\s*\d',
}


def select_categories(text: str) -> set[str]:
    lowered = text.lower()
    chosen = {name for name, pattern in CATEGORY_HINTS.items() if re.search(pattern, lowered)}
    if 'documents' in chosen or 'terminal' in chosen:
        chosen.add('files')
    return chosen
