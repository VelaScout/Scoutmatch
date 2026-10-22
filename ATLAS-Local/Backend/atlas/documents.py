"""Deterministic local document tools: validated input, exclusive publish, structural verification."""
from __future__ import annotations
import io
import os
from pathlib import Path
from xml.sax.saxutils import escape
from .tools import Result, ToolError


def destination(env, path, suffix):
    target = env.resolve(path, for_write=True)
    if target.suffix.lower() != suffix:
        raise ToolError(f'Dateiendung muss {suffix} sein.')
    if target.exists():
        raise ToolError('Datei existiert bereits. Bitte einen neuen Dateinamen wählen.')
    return target


def publish(target, data):
    target.parent.mkdir(parents=True, exist_ok=True)
    # O_EXCL also protects against a concurrent creator; never truncate another file.
    try:
        with target.open('xb') as output:
            output.write(data)
            output.flush()
            os.fsync(output.fileno())
    except FileExistsError as exc:
        raise ToolError('Datei existiert bereits. Bitte einen neuen Dateinamen wählen.') from exc
    except OSError:
        target.unlink(missing_ok=True)
        raise
    return Result(True, f'Erstellt: {target}', verified=True, data={'path': str(target), 'bytes': len(data)})


def text(value, label, limit):
    if not isinstance(value, str) or not value.strip() or len(value) > limit:
        raise ToolError(f'{label}: nichtleeren Text mit höchstens {limit} Zeichen angeben.')
    return value.strip()


def create_presentation(env, path: str, title: str, slides: list):
    target = destination(env, path, '.pptx')
    title = text(title, 'Titel', 120)
    if not isinstance(slides, list) or not 1 <= len(slides) <= 40:
        raise ToolError('Bitte 1 bis 40 Folien angeben.')
    clean = []
    for slide in slides:
        if not isinstance(slide, dict):
            raise ToolError('Jede Folie braucht title und bullets.')
        heading = text(slide.get('title'), 'Folientitel', 100)
        bullets = slide.get('bullets')
        if not isinstance(bullets, list) or not 1 <= len(bullets) <= 6:
            raise ToolError('Je Folie 1 bis 6 Stichpunkte angeben; längere Inhalte aufteilen.')
        clean.append((heading, [text(b, 'Stichpunkt', 180) for b in bullets]))
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
    except ImportError as exc:
        raise ToolError('PowerPoint-Modul fehlt. AKTUALISIEREN.command ausführen.') from exc
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    deck.core_properties.title = title
    for index, (heading, bullets) in enumerate(clean):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(245, 247, 251)
        def box(x, y, w, h):
            frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
            frame.word_wrap = True
            return frame
        frame = box(.7, .45, 11.9, .95)
        frame.text = heading
        for run in frame.paragraphs[0].runs:
            run.font.size, run.font.bold = Pt(30), True
            run.font.color.rgb = RGBColor(24, 45, 76)
        frame = box(.8, 1.65, 11.6, 5.1)
        for i, bullet in enumerate(bullets):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.text = '• ' + bullet
            p.font.size = Pt(20)
            p.space_after = Pt(12)
        footer = box(.8, 7.03, 11.7, .3)
        footer.text = f'{title[:75]}  |  {index + 1} / {len(clean)}'
        footer.paragraphs[0].font.size = Pt(10)
    buf = io.BytesIO()
    deck.save(buf)
    data = buf.getvalue()
    check = Presentation(io.BytesIO(data))
    if len(check.slides) != len(clean):
        raise ToolError('Präsentationsprüfung fehlgeschlagen.')
    return publish(target, data)


def create_pdf(env, path: str, title: str, content: str):
    target = destination(env, path, '.pdf')
    title = text(title, 'Titel', 160)
    content = text(content, 'Inhalt', 100_000)
    try:
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from pypdf import PdfReader
    except ImportError as exc:
        raise ToolError('PDF-Modul fehlt. AKTUALISIEREN.command ausführen.') from exc
    font = 'Helvetica'
    for candidate in ('/System/Library/Fonts/Supplemental/Arial.ttf',
                      '/Library/Fonts/Arial.ttf', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'):
        if Path(candidate).is_file():
            font = 'AtlasDocument'
            if font not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(font, candidate))
            break
    # Fail explicitly instead of silently replacing unsupported characters with black squares.
    charset = getattr(pdfmetrics.getFont(font).face, 'charToGlyph', None)
    if charset is not None:
        unsupported = {c for c in title + content if not c.isspace() and ord(c) not in charset}
        if unsupported:
            raise ToolError('PDF-Schrift unterstützt diese Zeichen nicht: ' + ''.join(sorted(unsupported))[:50])
    else:
        try: (title + content).encode('cp1252')
        except UnicodeEncodeError as exc:
            raise ToolError('Für diesen Text fehlt eine passende Unicode-Schrift.') from exc
    styles = getSampleStyleSheet()
    for name in ('Title', 'BodyText'):
        styles[name].fontName = font
    styles['BodyText'].fontSize = 11
    styles['BodyText'].leading = 16
    story = [Paragraph(escape(title), styles['Title']), Spacer(1, 16)]
    for paragraph in content.split('\n'):
        story.append(Paragraph(escape(paragraph) if paragraph.strip() else ' ', styles['BodyText']))
        story.append(Spacer(1, 5))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, title=title, author='ATLAS',
                      leftMargin=50, rightMargin=50, topMargin=50, bottomMargin=50).build(story)
    data = buf.getvalue()
    check = PdfReader(io.BytesIO(data))
    if not check.pages or not any(p.extract_text() for p in check.pages):
        raise ToolError('PDF-Prüfung fehlgeschlagen.')
    return publish(target, data)
