"""Real document bytes and agent regressions. LLM decisions are simulated."""
import tempfile
import threading
import unittest
from pathlib import Path
from atlas.agent import Agent, Request, Bridge, Limits, Action, task_context
from atlas.learning import Learning
from atlas.llm import Ollama, ModelError, route
from atlas.tools import Environment, build_tools, ToolError
from atlas.documents import create_pdf, create_presentation
from fake_ollama import FakeOllama, answer, tool_call

class ReleaseTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.home = Path(tmp.name)
        self.env = Environment(home=self.home)

    def test_pdf_real_text_and_multiple_pages(self):
        from pypdf import PdfReader
        result = create_pdf(self.env, 'Bericht.pdf', 'Überblick & Ziele', 'Größe: 42 € <Test>\n' * 160)
        self.assertTrue(result.verified)
        pdf = PdfReader(self.home / 'Bericht.pdf')
        self.assertGreater(len(pdf.pages), 1)
        text = ''.join(p.extract_text() for p in pdf.pages)
        self.assertIn('Überblick & Ziele', text)
        self.assertIn('Größe: 42 € <Test>', text)

    def test_powerpoint_real_slide_contents(self):
        from pptx import Presentation
        result = create_presentation(self.env, 'Plan.pptx', 'Trainingsplan', [
            {'title': 'Montag', 'bullets': ['Krafttraining', 'Erholung']},
            {'title': 'Dienstag', 'bullets': ['Fußball']}])
        deck = Presentation(self.home / 'Plan.pptx')
        self.assertTrue(result.verified)
        self.assertEqual(len(deck.slides), 2)
        self.assertIn('Fußball', '\n'.join(s.text for s in deck.slides[1].shapes if s.has_text_frame))

    def test_existing_files_and_wrong_suffixes_are_rejected(self):
        file = self.home / 'exists.pdf'
        file.write_bytes(b'original')
        with self.assertRaises(ToolError): create_pdf(self.env, str(file), 'Title', 'Body')
        self.assertEqual(file.read_bytes(), b'original')
        with self.assertRaises(ToolError): create_pdf(self.env, 'fake.txt', 'Title', 'Body')
        with self.assertRaises(ToolError): create_presentation(self.env, 'fake.pdf', 'Title', [])

    def test_outside_path_and_symlink_rejected(self):
        with self.assertRaises(ToolError): create_pdf(self.env, '/tmp/outside-atlas.pdf', 'Title', 'Body')
        (self.home / 'escape').symlink_to('/tmp', target_is_directory=True)
        with self.assertRaises(ToolError): create_pdf(self.env, 'escape/outside.pdf', 'Title', 'Body')

    def test_invalid_input_leaves_no_partial_file(self):
        for slides in ([], 'slides', [{}], [{'title': 'T', 'bullets': ['x'] * 7}],
                       [{'title': 'T', 'bullets': ['x' * 181]}]):
            with self.subTest(slides=slides):
                with self.assertRaises(ToolError): create_presentation(self.env, 'bad.pptx', 'Title', slides)
                self.assertFalse((self.home / 'bad.pptx').exists())

    def make_agent(self):
        self.fake = FakeOllama()
        self.addCleanup(self.fake.close)
        self.client = Ollama(base=self.fake.url)
        self.agent = Agent(self.client, build_tools(self.env), Bridge(), limits=Limits(seconds=10, max_rounds=4))

    def ask(self, text, history=None):
        return list(self.agent.run(Request(system='ATLAS', history=(history or []) + [
            {'role': 'user', 'content': text}]), threading.Event()))

    def test_followup_selects_previous_task_tools_and_executes(self):
        self.make_agent()
        self.fake.script = [tool_call('create_pdf', path='result.pdf', title='Plan', content='Test'), answer('PDF erstellt.')]
        events = self.ask('Mach das bitte', [{'role': 'user', 'content': 'Erstelle eine PDF mit meinem Plan'},
                                            {'role': 'assistant', 'content': 'Welchen Namen soll ich verwenden?'}])
        self.assertTrue((self.home / 'result.pdf').exists())
        names = [t['function']['name'] for t in self.fake.chat_requests()[0]['tools']]
        self.assertIn('create_pdf', names)
        self.assertEqual(events[-1]['mode'], 'agent')

    def test_new_topic_does_not_inherit_task(self):
        history = [{'role': 'user', 'content': 'Erstelle eine PDF'}, {'role': 'user', 'content': 'Warum ist der Himmel blau?'}]
        self.assertEqual(task_context(history, history[-1]['content']), history[-1]['content'])
        self.assertEqual(route('Eine PowerPoint Präsentation über Sport'), 'agent')

    def test_tool_batch_results_match_advertised_calls(self):
        self.make_agent()
        calls = [tool_call('create_folder', path=f'folder-{i}')['message']['tool_calls'][0] for i in range(6)]
        self.fake.script = [{'message': {'content': '', 'tool_calls': calls}}, answer('Vier Ordner erstellt.')]
        self.ask('Erstelle mehrere Ordner')
        messages = self.fake.chat_requests()[-1]['messages']
        requested = sum(len(m.get('tool_calls', [])) for m in messages)
        returned = sum(m['role'] == 'tool' for m in messages)
        self.assertEqual(requested, 6)
        self.assertEqual(returned, requested)
        self.assertFalse((self.home / 'folder-4').exists())

    def test_malformed_model_response_becomes_model_error(self):
        self.make_agent()
        for payload in ({'message': ['bad']}, {'message': 'bad'}, {'message': {'content': ['bad']}},
                        {'message': {'tool_calls': 'bad'}}, {'message': {'tool_calls': [None]}},
                        {'message': {'tool_calls': [{'function': {'name': 42}}]}}):
            with self.subTest(payload=payload):
                self.fake.script = [payload]
                with self.assertRaises(ModelError): self.client.chat('qwen3:4b', [{'role': 'user', 'content': 'Test'}])

    def test_partial_success_is_not_learned_as_complete(self):
        failed = Action('read_file', 'Read', 'missing.pdf', False, False, 'Missing')
        success = Action('read_file', 'Read', 'different.pdf', True, True, 'Content')
        self.assertEqual(Learning.evaluate([failed, success], True), ('teilweise', False))
        declined = Action('delete_path', 'Delete', 'file', False, False, 'No', declined=True)
        self.assertEqual(Learning.evaluate([success, declined], True), ('teilweise', False))
        fixed = Action('read_file', 'Read', 'missing.pdf', True, True, 'Content')
        self.assertEqual(Learning.evaluate([failed, fixed], True), ('erfolg', True))

    def test_mac_sandbox_temp_matches_child_environment_and_is_cleaned(self):
        import json
        import subprocess
        from unittest.mock import patch
        from atlas.skills import run_harness
        captured = {}
        def run(command, **kwargs):
            captured.update(kwargs)
            captured['command'] = command
            temp = Path(kwargs['env']['TMPDIR'])
            self.assertTrue(temp.is_dir())
            self.assertIn('(subpath ' + json.dumps(str(temp)) + ')', command[2])
            self.assertNotIn('(subpath "/private/var/folders")', command[2])
            self.assertEqual(kwargs['env']['TEMP'], str(temp))
            return subprocess.CompletedProcess(command, 0, '{"ok":true}\n', '')
        with patch('atlas.skills.shutil.which', return_value='/usr/bin/sandbox-exec'), patch('atlas.skills.subprocess.run', side_effect=run):
            self.assertTrue(run_harness('test', self.home, self.home)['ok'])
        self.assertFalse(Path(captured['env']['TMPDIR']).exists())

    def test_harness_can_create_temp_with_minimal_environment(self):
        from atlas.skills import run_harness
        skill = self.home / 'skill'
        output = self.home / 'output'
        skill.mkdir()
        (skill / 'skill.py').write_text('def run(args, output_dir):\n return {"ok": True}\n')
        result = run_harness('run', skill, output)
        self.assertTrue(result['ok'], result)
